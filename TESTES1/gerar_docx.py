from docx import Document
from docx.shared import Inches

# 1. Inicializa o construtor do documento
doc = Document()

# 2. Adiciona um título principal
doc.add_heading('Relatório Automatizado de Vendas', level=0)

# 3. Adiciona um parágrafo de texto normal
p = doc.add_paragraph('Este documento foi gerado automaticamente usando o script de ')
p.add_run('build docx').bold = True  # Adiciona texto em negrito
p.add_run(' em Python.')

# 4. Adiciona um subtítulo
doc.add_heading('Resultados do Trimestre', level=1)

# 5. Adiciona uma lista com marcadores (bullets)
doc.add_paragraph('Aumento de 15% nas vendas online.', style='List Bullet')
doc.add_paragraph('Expansão da equipe de suporte.', style='List Bullet')

# 6. Salva (constrói) o arquivo final
doc.save('Meu_Documento_Gerado.docx')

print("Arquivo DOCX construído com sucesso!")